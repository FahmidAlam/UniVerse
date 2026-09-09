"""The rendered workbook must carry nothing from the term the template came from.

The template is a clone of the real Summer-2025 routine and it is not clean. It
holds classes, faculty acronyms, coloured annotation and hand-drawn separators
that no part of the pipeline regenerates. `cell.value = None` does not remove
formatting, so all of that used to survive into every generated routine:

  * 27 cells inside the cohort grid carry an explicit RED font and 3 carry a
    green or yellow fill. Whatever class the solver placed there came out
    coloured, appearing to mean something it never meant.
  * Stale THICK top borders on rows 18, 27, 32, 34, 41, 46, 52, 57 and 59 drew
    heavy separators in the middle of a batch, because Excel draws the line
    between two rows as the union of the upper bottom and the lower top.
  * A frozen "66 Batch Section Distribution" student-ID table and a per-day
    teacher-on-duty roster sat below the grid.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from conftest import build_workbook, make_config, make_row

import ingest
import render
import solver

DAYS = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
        "Saturday"]


def style(side) -> str | None:
    """A cleared border side is `None`, not a Side with `style=None`."""
    return side.style if side is not None else None


def own_border_cols(ws, row):
    """Grid columns whose border this row actually controls.

    A cell swallowed by a merge (Friday's two-column BREAK block covers column
    H) cannot carry its own border, so it is not part of the separator.
    """
    from openpyxl.cell.cell import MergedCell
    return [c for c in render.GRID_COLS
            if not isinstance(ws.cell(row, c), MergedCell)]


def build(spec, **cfg_kw):
    rows = [make_row(batch=b, section=s, code=c, teacher=t, classes=28)
            for b, sections in spec.items()
            for s in sections
            for c, t in (("CSE-1101", "T1"), ("CSE-1103", "T2"))]
    cfg = solver.load_config(
        override=make_config(teachers=["T1", "T2"], **cfg_kw))
    ds = ingest.ingest_bytes(build_workbook(rows),
                             weeks_in_term=cfg["weeks_in_term"])
    res = solver.solve(ds, cfg, time_limit_s=10)
    data = render.render_bytes(res["rows"], ds["cohorts"], cfg)
    return ds, openpyxl.load_workbook(io.BytesIO(data))


@pytest.fixture(scope="module")
def rendered():
    return build({"70": ["A", "B", "C"], "69": ["A", "B"], "68": ["A"]})


# ── colour ───────────────────────────────────────────────────────────────

def test_the_template_really_does_carry_the_red(the_template=None):
    """Guard the diagnosis: if a future template no longer has these, the
    normalisation below is proving nothing."""
    wb = openpyxl.load_workbook(render.TEMPLATE)
    reds = 0
    for day in DAYS:
        ws = wb[day]
        for r in range(render.DATA_FIRST_ROW, 84):
            for c in render.GRID_COLS:
                f = ws.cell(r, c).font
                rgb = f.color.rgb if (f.color and f.color.type == "rgb") else None
                if isinstance(rgb, str) and rgb.upper() == "FFFF0000":
                    reds += 1
    assert reds > 0, "template no longer has red cells; revisit this test"


def test_no_class_is_ever_rendered_in_a_colour(rendered):
    _, wb = rendered
    bad = []
    for day in DAYS:
        ws = wb[day]
        for r in range(render.DATA_FIRST_ROW, render._data_last_row(ws) + 1):
            for c in render.GRID_COLS:
                cell = ws.cell(r, c)
                f = cell.font
                rgb = f.color.rgb if (f.color and f.color.type == "rgb") else None
                if isinstance(rgb, str) and rgb.upper() not in ("FF000000",
                                                                "00000000"):
                    bad.append((day, cell.coordinate, "font", rgb))
                fill = cell.fill
                frgb = (fill.fgColor.rgb
                        if (fill.patternType == "solid" and fill.fgColor
                            and fill.fgColor.type == "rgb") else None)
                if isinstance(frgb, str) and frgb.upper() not in ("FFFFFFFF",
                                                                  "00000000"):
                    bad.append((day, cell.coordinate, "fill", frgb))
    assert bad == []


# ── borders ──────────────────────────────────────────────────────────────

def test_the_heavy_rule_falls_only_at_the_end_of_a_batch(rendered):
    ds, wb = rendered
    groups = render.cohort_groups(ds["cohorts"])
    last_rows = {render.DATA_FIRST_ROW + g["last"] for g in groups}
    for day in DAYS:
        ws = wb[day]
        for i in range(len(ds["cohorts"])):
            r = render.DATA_FIRST_ROW + i
            styles = {style(ws.cell(r, c).border.bottom)
                      for c in own_border_cols(ws, r)}
            want = "thick" if r in last_rows else "thin"
            assert styles == {want}, f"{day} row {r} ({ds['cohorts'][i]})"


def test_a_rows_top_always_agrees_with_the_row_above_it(rendered):
    """The bug: Excel unions the two, so a stale thick top drew a separator the
    bottom never asked for."""
    ds, wb = rendered
    n = len(ds["cohorts"])
    for day in DAYS:
        ws = wb[day]
        for i in range(1, n):
            r = render.DATA_FIRST_ROW + i
            shared = set(own_border_cols(ws, r - 1)) &                 set(own_border_cols(ws, r))
            for c in sorted(shared):
                above = style(ws.cell(r - 1, c).border.bottom)
                here = style(ws.cell(r, c).border.top)
                assert above == here, f"{day} row {r} col {c}"


def test_the_table_is_closed_top_and_bottom(rendered):
    ds, wb = rendered
    n = len(ds["cohorts"])
    for day in DAYS:
        ws = wb[day]
        first, last = render.DATA_FIRST_ROW, render.DATA_FIRST_ROW + n - 1
        assert {style(ws.cell(first, c).border.top)
                for c in own_border_cols(ws, first)} == {"thick"}
        assert {style(ws.cell(last, c).border.bottom)
                for c in own_border_cols(ws, last)} == {"thick"}


def test_unused_rows_keep_no_separator_from_a_longer_previous_routine(rendered):
    ds, wb = rendered
    n = len(ds["cohorts"])
    for day in DAYS:
        ws = wb[day]
        for r in range(render.DATA_FIRST_ROW + n,
                       render._data_last_row(ws) + 1):
            for c in own_border_cols(ws, r):
                b = ws.cell(r, c).border
                assert style(b.bottom) is None, f"{day} row {r}"
                assert style(b.top) is None, f"{day} row {r}"


def test_one_section_per_batch_makes_every_line_heavy():
    ds, wb = build({"70": ["A"], "69": ["A"], "68": ["A"]})
    ws = wb["Sunday"]
    for i in range(3):
        r = render.DATA_FIRST_ROW + i
        assert {style(ws.cell(r, c).border.bottom)
                for c in own_border_cols(ws, r)} == {"thick"}


def test_one_batch_of_many_sections_has_a_single_heavy_line_at_the_end():
    ds, wb = build({"72": list("ABCDEF")})
    ws = wb["Sunday"]
    styles = [
        {style(ws.cell(render.DATA_FIRST_ROW + i, c).border.bottom)
         for c in own_border_cols(ws, render.DATA_FIRST_ROW + i)}.pop()
        for i in range(6)
    ]
    assert styles == ["thin"] * 5 + ["thick"]


# ── stale term content ───────────────────────────────────────────────────

def test_the_frozen_batch_section_distribution_table_is_gone(rendered):
    _, wb = rendered
    for day in DAYS:
        ws = wb[day]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                if isinstance(cell.value, str):
                    assert "batch section distribution" not in \
                        cell.value.strip().lower(), f"{day} {cell.coordinate}"


def test_the_old_student_id_ranges_are_gone(rendered):
    _, wb = rendered
    for day in DAYS:
        ws = wb[day]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                if isinstance(cell.value, str) and "..........." in cell.value:
                    pytest.fail(f"{day} {cell.coordinate}: {cell.value!r}")


def test_the_stale_teacher_on_duty_roster_is_gone(rendered):
    """The row under STUDENT NO. held real faculty acronyms from a term that
    has ended. The label column stays; the names do not."""
    _, wb = rendered
    for day in DAYS:
        ws = wb[day]
        duty = None
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if isinstance(v, str) and v.strip().lower().startswith("student no"):
                duty = r + 1
                break
        assert duty is not None, f"{day}: furniture not found"
        values = [ws.cell(duty, c).value for c in range(3, 12)]
        assert all(v is None for v in values), f"{day} row {duty}: {values}"


def test_page_furniture_itself_still_survives(rendered):
    """Bus timings and headcount labels are campus furniture, not routine data,
    and must not be swept away with the stale blocks."""
    _, wb = rendered
    for day in DAYS:
        ws = wb[day]
        labels = {str(ws.cell(r, 2).value).strip().lower()
                  for r in range(1, ws.max_row + 1)
                  if isinstance(ws.cell(r, 2).value, str)}
        assert "bus time" in labels, day
        assert any(x.startswith("student no") for x in labels), day


# ── the break column follows configuration, not the day's name ───────────

def _one_friday(**cfg_kw):
    """One cohort, four courses, Friday only, so every usable period is filled
    and period 4 (column H, beside the break) is exercised."""
    rows = [make_row(batch="70", section="A", code=c, teacher=t, classes=14)
            for c, t in (("CSE-1101", "T1"), ("CSE-1103", "T2"),
                         ("CSE-1105", "T3"), ("CSE-1107", "T4"))]
    cfg = solver.load_config(override=make_config(
        teachers=["T1", "T2", "T3", "T4"], days=["Friday"], **cfg_kw))
    ds = ingest.ingest_bytes(build_workbook(rows),
                             weeks_in_term=cfg["weeks_in_term"])
    res = solver.solve(ds, cfg, time_limit_s=10)
    wb = openpyxl.load_workbook(
        io.BytesIO(render.render_bytes(res["rows"], ds["cohorts"], cfg)))
    return cfg, res, wb["Friday"]


def test_a_taught_friday_p4_renders_instead_of_crashing_the_job():
    """The regression this replaces did not draw the wrong thing — it aborted
    the whole generation with `'MergedCell' object attribute 'value' is
    read-only`, because the template pre-merges Friday's G:H and the renderer
    assumed period 4 was never taught there."""
    cfg, res, ws = _one_friday(friday_no_p4=False)
    assert render._break_span("Friday", cfg) == 1
    printed = {str(ws.cell(3, c).value) for c in range(4, 12)
               if ws.cell(3, c).value}
    for r in res["rows"]:
        assert any(r["subject_code"] in v for v in printed), r["subject_code"]


def test_a_taught_friday_p4_lands_in_its_own_column():
    cfg, res, ws = _one_friday(friday_no_p4=False)
    p4 = next(r for r in res["rows"] if r["period"] == 4)
    assert p4["subject_code"] in str(ws.cell(3, render.BREAK_COL + 1).value)
    assert str(ws.cell(3, render.BREAK_COL).value) in ("B", "R", "E", "A", "K")


def test_a_blocked_friday_p4_still_widens_the_break_across_both_columns():
    rows = [make_row(batch="70", section="A", code=c, teacher=t, classes=14)
            for c, t in (("CSE-1101", "T1"), ("CSE-1103", "T2"),
                         ("CSE-1105", "T3"))]
    cfg = solver.load_config(override=make_config(
        teachers=["T1", "T2", "T3"], days=["Friday"], friday_no_p4=True))
    ds = ingest.ingest_bytes(build_workbook(rows),
                             weeks_in_term=cfg["weeks_in_term"])
    res = solver.solve(ds, cfg, time_limit_s=10)
    ws = openpyxl.load_workbook(
        io.BytesIO(render.render_bytes(res["rows"], ds["cohorts"], cfg)))["Friday"]
    assert render._break_span("Friday", cfg) == 2
    assert ws.cell(3, render.BREAK_COL).value == "BREAK"
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert any(m.startswith("G3:H") for m in merged), merged


def test_a_generic_blocked_period_widens_the_break_on_that_day_too():
    """Nothing about this is Friday-specific any more: block period 4 on
    Tuesday and Tuesday's break widens, while Friday's does not."""
    cfg = solver.load_config(override=make_config(
        friday_no_p4=False, blocked_periods={"Tuesday": [4]}))
    assert render._break_span("Tuesday", cfg) == 2
    assert render._break_span("Friday", cfg) == 1
