"""Workbook rendering.

The renderer clones a template that carries more cohort rows than any run
writes, plus page furniture below them. These tests pin the boundary between
the two, because getting it wrong put a phantom class into every printed
routine and double-booked a room.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from conftest import build_workbook, make_config, make_row

import ingest
import render
import solver

DAYS = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
        "Saturday"]
#: Sheet column -> period index, mirroring the template's `col` letters
#: (G, column 7, is the break column).
PCOL = {4: 1, 5: 2, 6: 3, 8: 4, 9: 5, 10: 6, 11: 7}


def render_fixture(rows=None, **cfg_kw):
    rows = rows or [make_row(section=s, code=c, teacher=t)
                    for s in ("A", "B")
                    for c, t in (("CSE-1101", "AAA"), ("CSE-1102", "AAA"),
                                 ("CSE-1151", "BBB"))]
    cfg = solver.load_config(override=make_config(**cfg_kw))
    ds = ingest.ingest_bytes(build_workbook(rows),
                             weeks_in_term=cfg["weeks_in_term"])
    res = solver.solve(ds, cfg, time_limit_s=8)
    data = render.render_bytes(res["rows"], ds["cohorts"], cfg)
    return res, openpyxl.load_workbook(io.BytesIO(data))


def class_cells(wb):
    """Every non-empty class cell inside the cohort block.

    Bounded by `_data_last_row` deliberately: the page furniture underneath
    (bus timings, headcounts) also writes into the period columns and is not
    part of the timetable.
    """
    out = []
    for d in DAYS:
        ws = wb[d]
        for r in range(render.DATA_FIRST_ROW, render._data_last_row(ws) + 1):
            for col, p in PCOL.items():
                v = ws.cell(row=r, column=col).value
                if v and str(v).strip() not in ("B", "R"):
                    out.append((d, r, p, " ".join(str(v).split())))
    return out


def written_text(res):
    return {f"{r['subject_code']} {r['teacher_code']} {r['room']}"
            for r in res["rows"]}


@pytest.fixture(scope="module")
def rendered():
    """One render shared by the read-only tests — each solve costs seconds."""
    return render_fixture()


# ── the template must not leak its own content ───────────────────────────

def test_no_class_from_the_template_survives_a_render(rendered):
    """The shipped template carries hard-coded ACM classes on Friday rows 59
    and 83, left over from the workbook it was cloned from. DATA_LAST_ROW was
    82, so row 59 was cleared and row 83 was not: every rendered Friday printed
    a phantom `ACM-2000 *** GL`, double-booking a room the solver had already
    given to a real class.
    """
    res, wb = rendered
    written = written_text(res)
    leaked = [c for c in class_cells(wb) if c[3] not in written]
    assert leaked == [], f"template content survived the render: {leaked}"


def test_the_specific_acm_rows_are_gone(rendered):
    _res, wb = rendered
    text = " ".join(c[3] for c in class_cells(wb))
    assert "ACM-1000" not in text
    assert "ACM-2000" not in text
    assert "***" not in text, "placeholder teacher leaked from the template"


def test_page_furniture_below_the_cohorts_is_preserved(rendered):
    """Bus timings, headcounts and the teacher-on-duty row are part of the
    printed page and must survive — the fix clears cohort space, not the whole
    sheet."""
    _res, wb = rendered
    ws = wb["Friday"]
    labels = {str(ws.cell(row=r, column=1).value or
                  ws.cell(row=r, column=2).value or "").strip().lower()
              for r in range(render.DATA_FIRST_ROW, (ws.max_row or 0) + 1)}
    assert "bus time" in labels
    assert "student no." in labels


def test_data_block_boundary_is_found_not_assumed(rendered):
    """Each sheet's cohort block ends exactly where its page furniture starts.

    The boundary is not the same on every sheet — Tuesday and Wednesday stop at
    58, Friday runs to 83 — which is precisely why one DATA_LAST_ROW constant
    could never be right for all of them.
    """
    _res, wb = rendered
    for d in DAYS:
        ws = wb[d]
        last = render._data_last_row(ws)
        below = str(ws.cell(row=last + 1, column=1).value or
                    ws.cell(row=last + 1, column=2).value or "").strip().lower()
        assert below in render._FURNITURE_LABELS, \
            f"{d}: row after the block is {below!r}, not page furniture"


def test_friday_block_covers_the_row_the_old_constant_missed(rendered):
    """Friday is the sheet that broke: its cohort block runs to row 83, one
    past the old DATA_LAST_ROW of 82, which is where `ACM-2000 *** GL` sat."""
    _res, wb = rendered
    ws = wb["Friday"]
    assert render._data_last_row(ws) == 83
    assert all(ws.cell(row=83, column=col).value in (None, "") for col in PCOL)


# ── every scheduled class reaches the page, once ─────────────────────────

def test_every_solver_row_is_rendered_exactly_once(rendered):
    res, wb = rendered
    cells = class_cells(wb)
    assert len(cells) == len(res["rows"])
    assert len({(d, r, p) for d, r, p, _ in cells}) == len(cells)


def test_a_room_is_never_printed_twice_in_one_slot(rendered):
    """The end-to-end version of the leak: whatever appears on the page must
    not double-book a room, whether it came from the solver or the template."""
    _res, wb = rendered
    seen = {}
    for d, _r, p, text in class_cells(wb):
        room = text.split(" ")[-1]
        key = (room, d, p)
        assert key not in seen, \
            f"{room} used twice on {d} P{p}: {seen[key]} / {text}"
        seen[key] = text


@pytest.mark.parametrize("day", DAYS)
def test_no_sheet_renders_a_class_for_another_day(day, rendered):
    res, wb = rendered
    on_page = {(d, t) for d, _r, _p, t in class_cells(wb)}
    for r in res["rows"]:
        if r["day"] == day:
            text = f"{r['subject_code']} {r['teacher_code']} {r['room']}"
            assert (day, text) in on_page


# ── headers and blocked periods follow the configuration ─────────────────

def test_period_headers_come_from_the_configured_times():
    _res, wb = render_fixture(periods=[
        {"idx": 1, "label": "", "start": "09:30", "end": "10:45", "col": "D"},
        {"idx": 2, "label": "", "start": "10:45", "end": "12:00", "col": "E"},
    ])
    header = str(wb["Sunday"].cell(row=render.HEADER_ROW, column=4).value)
    assert "09:30" in header and "10:45" in header


def test_a_blocked_period_is_left_blank_on_that_day_only():
    _res, wb = render_fixture(blocked_periods={"Friday": [4]})
    friday = [(p, t) for d, _r, p, t in class_cells(wb) if d == "Friday"]
    assert all(p != 4 for p, _ in friday), friday
    # The column still exists for the other days.
    assert wb["Sunday"].cell(row=render.HEADER_ROW, column=8).value is not None
