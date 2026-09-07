"""Normalization rules: raw distribution rows -> academic offerings.

These pin the two behaviours the engine used to get wrong — one row is not one
course, and one course is not automatically two weekly classes.
"""

from __future__ import annotations

import pytest
from conftest import build_workbook, make_row

import ingest


def norm(rows, weeks=14):
    return ingest.ingest_bytes(build_workbook(rows), weeks_in_term=weeks)


def offering(ds, cohort, code):
    for o in ds["offerings"]:
        if o["cohort"] == cohort and o["code"] == code:
            return o
    raise AssertionError(f"{code} missing for {cohort}; "
                         f"got {[(o['cohort'], o['code']) for o in ds['offerings']]}")


# ── credit / class-count -> weekly sessions ──────────────────────────────

@pytest.mark.parametrize("classes,weeks,expected", [
    (28, 14, 2),   # 3.0 cr theory and 1.5 cr lab in the real file
    (19, 14, 2),   # 2.0 cr and 1.0 cr — ceil(1.36) == 2
    (38, 14, 3),   # CSE-4116: the case the old engine under-scheduled
    (14, 14, 1),   # once a week
    (42, 14, 3),
    (28, 12, 3),   # shorter term => same course meets more often
])
def test_required_sessions_follow_the_class_count(classes, weeks, expected):
    ds = norm([make_row(classes=classes)], weeks=weeks)
    o = offering(ds, "66-A", "CSE-1101")
    assert o["required_sessions"] == expected
    assert o["sessions_basis"] == "classes_total"
    assert len([s for s in ds["sessions"] if s["oid"] == o["oid"]]) == expected


def test_class_per_week_is_only_a_fallback():
    """`Class/Week` is a constant 2 in the real file even for courses that meet
    three times, so it must never override an explicit class count."""
    ds = norm([make_row(classes=38, per_week=2)])
    assert offering(ds, "66-A", "CSE-1101")["required_sessions"] == 3

    ds = norm([make_row(classes=None, per_week=3)])
    o = offering(ds, "66-A", "CSE-1101")
    assert o["required_sessions"] == 3
    assert o["sessions_basis"] == "class_per_week"


def test_missing_counts_fall_back_to_default_and_warn():
    ds = norm([make_row(classes=None, per_week=None)])
    o = offering(ds, "66-A", "CSE-1101")
    assert o["required_sessions"] == ingest.DEFAULT_SESSIONS_PER_WEEK
    assert o["sessions_basis"] == "default"
    assert any("assuming" in w for w in ds["warnings"])


# ── team teaching ────────────────────────────────────────────────────────

def test_split_teacher_rows_collapse_into_one_offering():
    """The defect that inflated the Summer-2025 routine by 36 classes: two
    rows, 14 classes each, is ONE course meeting twice — not four times."""
    ds = norm([
        make_row(code="CSE-2111", teacher="AAA", classes=14),
        make_row(code="CSE-2111", teacher="BBB", classes=14),
    ])
    assert len(ds["offerings"]) == 1
    o = offering(ds, "66-A", "CSE-2111")
    assert o["classes_total"] == 28
    assert o["required_sessions"] == 2
    assert len(ds["sessions"]) == 2


def test_first_row_owns_the_printed_slot_and_co_teachers_are_kept():
    ds = norm([
        make_row(code="CSE-2111", teacher="AAA", classes=14),
        make_row(code="CSE-2111", teacher="BBB", classes=14),
    ])
    o = offering(ds, "66-A", "CSE-2111")
    assert o["teachers"] == ["AAA", "BBB"]
    assert all(s["teacher"] == "AAA" for s in ds["sessions"])
    assert all(s["teachers"] == ["AAA", "BBB"] for s in ds["sessions"])
    assert ds["meta"]["shared_offerings"] == 1
    assert any("shared by AAA, BBB" in w for w in ds["warnings"])


def test_uneven_split_still_sums_to_the_term_total():
    """66-D CSE-1101 in the real file: 9 + 10 = 19 classes."""
    ds = norm([
        make_row(code="CSE-1101", teacher="AAA", classes=9, credit=2.0),
        make_row(code="CSE-1101", teacher="BBB", classes=10, credit=2.0),
    ])
    o = offering(ds, "66-A", "CSE-1101")
    assert o["classes_total"] == 19
    assert o["required_sessions"] == 2


def test_same_course_in_different_cohorts_stays_separate():
    ds = norm([
        make_row(section="A", code="CSE-1101"),
        make_row(section="B", code="CSE-1101"),
    ])
    assert len(ds["offerings"]) == 2
    assert {o["cohort"] for o in ds["offerings"]} == {"66-A", "66-B"}


# ── nothing is dropped silently ──────────────────────────────────────────

def test_every_exclusion_is_typed_and_traceable():
    ds = norm([
        make_row(code="CSE-1101"),
        make_row(code="CSE-3240", teacher=None),     # project: no teacher
        make_row(batch=None, code="ACM-1000"),       # not a cohort offering
    ])
    assert len(ds["offerings"]) == 1
    reasons = {e["reason"] for e in ds["excluded"]}
    assert reasons == {ingest.EXCLUDE_NO_TEACHER, ingest.EXCLUDE_NO_BATCH}
    for e in ds["excluded"]:
        assert e["row"] >= ingest.DATA_START
        assert e["detail"]
    assert ds["meta"]["excluded"] == 2


def test_blank_spacer_rows_are_not_reported_as_exclusions():
    rows = [make_row(code="CSE-1101"),
            make_row(batch=None, section=None, code=None, teacher=None,
                     classes=None, credit=None, duration=None, per_week=None),
            make_row(code="CSE-1151")]
    ds = norm(rows)
    assert len(ds["offerings"]) == 2
    assert ds["excluded"] == []


# ── classification ───────────────────────────────────────────────────────

def test_even_final_digit_marks_a_lab():
    ds = norm([make_row(code="CSE-1101"), make_row(code="CSE-1102")])
    assert offering(ds, "66-A", "CSE-1101")["is_lab"] is False
    assert offering(ds, "66-A", "CSE-1102")["is_lab"] is True


def test_service_offerings_are_flagged_and_excluded_from_cohorts():
    ds = norm([make_row(code="CSE-1101"),
               make_row(batch="BuA", section="GED A", code="GED-1122")])
    assert offering(ds, "BuA-GED A", "GED-1122")["is_service"] is True
    assert ds["cohorts"] == ["66-A"]
    assert ds["meta"]["service_sessions"] == 2


def test_cohorts_render_newest_batch_first():
    ds = norm([make_row(batch="64", section="B"), make_row(batch="66", section="A"),
               make_row(batch="64", section="A")])
    assert ds["cohorts"] == ["66-A", "64-A", "64-B"]


# ── the parser is bound to header text, not column order ─────────────────

def test_columns_may_be_reordered(monkeypatch):
    import conftest
    shuffled = ["Course Code", "Batch", "Teacher", "Section", "No. Of Classes",
                "Course Title", "Credit", "Prerequisite",
                "Conducting Department", "No. Of Students", "Class Duration",
                "Class/Week"]
    monkeypatch.setattr(conftest, "HEADERS", shuffled)

    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Course Distribution"
    for c, h in enumerate(shuffled, start=1):
        ws.cell(row=5, column=c, value=h)
    ws.cell(row=6, column=1, value="CSE-1101")
    ws.cell(row=6, column=2, value="66")
    ws.cell(row=6, column=3, value="AAA")
    ws.cell(row=6, column=4, value="A")
    ws.cell(row=6, column=5, value=38)
    buf = io.BytesIO()
    wb.save(buf)

    ds = ingest.ingest_bytes(buf.getvalue(), weeks_in_term=14)
    o = offering(ds, "66-A", "CSE-1101")
    assert o["teachers"] == ["AAA"]
    assert o["required_sessions"] == 3


def test_missing_required_column_fails_loudly():
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Course Distribution"
    for c, h in enumerate(["Batch", "Section", "Course Title"], start=1):
        ws.cell(row=5, column=c, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Course Code"):
        ingest.ingest_bytes(buf.getvalue())


def test_wrong_sheet_name_fails_loudly():
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Course Distribution"):
        ingest.ingest_bytes(buf.getvalue())
